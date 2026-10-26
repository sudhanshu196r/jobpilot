"""ATS-friendly PDF and DOCX output.

Layout rules that keep applicant-tracking parsers happy: one column, real
text (no images or tables), standard section headings, simple bullets, and
contact details in the body rather than a page header.
"""

from __future__ import annotations

import os
from pathlib import Path
from xml.sax.saxutils import escape

from ..schemas import ResumeDoc
from .profile_utils import display_month

_FONT_CANDIDATES = [
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("LiberationSans", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("Arial", "/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("Arial", "/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("Arial", "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]
_fonts: tuple[str, str] | None = None

_ASCII = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-",
                        "\u2022": "-", "\u2026": "...", "\u20b9": "Rs.", "\u00a0": " "})


def _register_fonts() -> tuple[str, str]:
    """A Unicode TTF if one exists on this machine, else built-in Helvetica."""
    global _fonts
    if _fonts:
        return _fonts
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont(f"JP-{name}", regular))
                pdfmetrics.registerFont(TTFont(f"JP-{name}-Bold", bold))
                _fonts = (f"JP-{name}", f"JP-{name}-Bold")
                return _fonts
            except Exception:  # pragma: no cover - corrupt font file
                continue
    _fonts = ("Helvetica", "Helvetica-Bold")
    return _fonts


def _date_range(start: str, end: str, current: bool) -> str:
    s = display_month(start) if start else ""
    e = "Present" if current or (start and not end) else (display_month(end) if end else "")
    return f"{s} – {e}" if s and e else (s or e)


def _contact_line(doc: ResumeDoc) -> list[str]:
    c = doc.contact
    return [x for x in (c.location, c.phone, c.email, c.linkedin, c.github, c.website) if x]


def render_pdf(doc: ResumeDoc, path: Path) -> Path:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    regular, bold = _register_fonts()
    unicode_ok = regular != "Helvetica"

    def t(value: str) -> str:
        value = value if unicode_ok else value.translate(_ASCII).encode("latin-1", "replace").decode("latin-1")
        return escape(value)

    ink, muted, accent = HexColor("#111111"), HexColor("#444444"), HexColor("#1f3a5f")
    styles = {
        "name": ParagraphStyle("name", fontName=bold, fontSize=19, leading=23, textColor=ink, alignment=TA_LEFT),
        "headline": ParagraphStyle("headline", fontName=regular, fontSize=10.5, leading=14, textColor=accent),
        "contact": ParagraphStyle("contact", fontName=regular, fontSize=9, leading=12, textColor=muted),
        "h": ParagraphStyle("h", fontName=bold, fontSize=10.5, leading=13, textColor=accent, spaceBefore=9, spaceAfter=2),
        "body": ParagraphStyle("body", fontName=regular, fontSize=9.6, leading=12.6, textColor=ink),
        "role": ParagraphStyle("role", fontName=bold, fontSize=10, leading=13, textColor=ink, spaceBefore=4),
        "meta": ParagraphStyle("meta", fontName=regular, fontSize=9, leading=12, textColor=muted),
        "bullet": ParagraphStyle("bullet", fontName=regular, fontSize=9.6, leading=12.4, textColor=ink),
    }

    story: list = []
    story.append(Paragraph(t(doc.contact.name or "Your Name"), styles["name"]))
    if doc.headline:
        story.append(Paragraph(t(doc.headline), styles["headline"]))
    if contact := _contact_line(doc):
        story.append(Paragraph("  |  ".join(t(x) for x in contact), styles["contact"]))

    def heading(label: str):
        story.append(Paragraph(label.upper(), styles["h"]))
        story.append(HRFlowable(width="100%", thickness=0.6, color=accent, spaceBefore=0, spaceAfter=3))

    def bullets(items: list[str]):
        if not items:
            return None
        return ListFlowable(
            [ListItem(Paragraph(t(b), styles["bullet"]), leftIndent=10, value="•") for b in items],
            bulletType="bullet", start="•", leftIndent=10, bulletFontName=regular, bulletFontSize=8,
            spaceBefore=1, spaceAfter=1,
        )

    if doc.summary:
        heading("Summary")
        story.append(Paragraph(t(doc.summary), styles["body"]))
    if doc.skills:
        heading("Skills")
        story.append(Paragraph(t(", ".join(doc.skills)), styles["body"]))
    if doc.experience:
        heading("Experience")
        for e in doc.experience:
            right = _date_range(e.start, e.end, e.current)
            title = f"<b>{t(e.title)}</b>" + (f" — {t(e.company)}" if e.company else "")
            meta = "  |  ".join(t(x) for x in (right, e.location) if x)
            block = [Paragraph(title, styles["role"])]
            if meta:
                block.append(Paragraph(meta, styles["meta"]))
            bl = bullets(e.bullets)
            if bl:
                block.append(bl)
            story.append(KeepTogether(block[:3]))
    if doc.projects:
        heading("Projects")
        for p in doc.projects:
            line = f"<b>{t(p.name)}</b>" + (f"  |  {t(p.url)}" if p.url else "")
            block = [Paragraph(line, styles["role"])]
            if p.description:
                block.append(Paragraph(t(p.description), styles["body"]))
            bl = bullets(p.bullets)
            if bl:
                block.append(bl)
            story.append(KeepTogether(block))
    if doc.education:
        heading("Education")
        for ed in doc.education:
            line = f"<b>{t(ed.degree)}</b>" + (f" — {t(ed.institution)}" if ed.institution else "")
            story.append(Paragraph(line, styles["role"]))
            meta = "  |  ".join(t(x) for x in (_date_range(ed.start, ed.end, False) if ed.start else display_month(ed.end),
                                                ed.location, ed.grade) if x)
            if meta:
                story.append(Paragraph(meta, styles["meta"]))
    if doc.certifications:
        heading("Certifications")
        story.append(bullets([" — ".join(x for x in (c.name, c.issuer, c.year) if x) for c in doc.certifications]))
    if doc.languages:
        heading("Languages")
        story.append(Paragraph(t(", ".join(doc.languages)), styles["body"]))

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"{doc.contact.name} - Resume", author=doc.contact.name, subject=doc.headline,
    )
    story.append(Spacer(1, 1))
    pdf.build(story)
    return path


def render_docx(doc: ResumeDoc, path: Path) -> Path:
    import docx
    from docx.enum.text import WD_BREAK  # noqa: F401
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    for section in d.sections:
        section.left_margin = section.right_margin = Pt(46)
        section.top_margin = section.bottom_margin = Pt(40)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    accent = RGBColor(0x1F, 0x3A, 0x5F)

    def para(text: str = "", bold_text: bool = False, size: float | None = None, color=None, space_after: float = 2):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        if text:
            r = p.add_run(text)
            r.bold = bold_text
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
        return p

    def heading(label: str):
        p = para(label.upper(), True, 11, accent, 2)
        p.paragraph_format.space_before = Pt(8)

    def bullet(text: str):
        p = d.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    para(doc.contact.name or "Your Name", True, 18)
    if doc.headline:
        para(doc.headline, False, 11, accent)
    if contact := _contact_line(doc):
        para("  |  ".join(contact), False, 9.5)
    if doc.summary:
        heading("Summary")
        para(doc.summary)
    if doc.skills:
        heading("Skills")
        para(", ".join(doc.skills))
    if doc.experience:
        heading("Experience")
        for e in doc.experience:
            p = para(space_after=0)
            p.add_run(e.title).bold = True
            if e.company:
                p.add_run(f" — {e.company}")
            meta = "  |  ".join(x for x in (_date_range(e.start, e.end, e.current), e.location) if x)
            if meta:
                para(meta, False, 9.5)
            for b in e.bullets:
                bullet(b)
    if doc.projects:
        heading("Projects")
        for pr in doc.projects:
            p = para(space_after=0)
            p.add_run(pr.name).bold = True
            if pr.url:
                p.add_run(f"  |  {pr.url}")
            if pr.description:
                para(pr.description)
            for b in pr.bullets:
                bullet(b)
    if doc.education:
        heading("Education")
        for ed in doc.education:
            p = para(space_after=0)
            p.add_run(ed.degree).bold = True
            if ed.institution:
                p.add_run(f" — {ed.institution}")
            meta = "  |  ".join(x for x in (_date_range(ed.start, ed.end, False) if ed.start else display_month(ed.end),
                                             ed.location, ed.grade) if x)
            if meta:
                para(meta, False, 9.5)
    if doc.certifications:
        heading("Certifications")
        for c in doc.certifications:
            bullet(" — ".join(x for x in (c.name, c.issuer, c.year) if x))
    if doc.languages:
        heading("Languages")
        para(", ".join(doc.languages))

    path.parent.mkdir(parents=True, exist_ok=True)
    d.core_properties.author = doc.contact.name
    d.core_properties.title = f"{doc.contact.name} - Resume"
    d.save(str(path))
    return path
